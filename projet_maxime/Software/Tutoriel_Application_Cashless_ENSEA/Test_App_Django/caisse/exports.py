"""
Construction du classeur Excel d'export des ventes d'un pole.

Regle de confidentialite : l'acheteur n'est jamais nomme. Chaque ligne ne
porte qu'un identifiant technique (ETU-00042), jamais le nom ni l'email,
pour qu'on ne puisse pas relier une ligne a une personne a la simple
lecture du fichier.
"""

from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import LigneTransaction, Transaction

_POLICE_ENTETE = Font(color="FFFFFF", bold=True)
_REMPLISSAGE_ENTETE = PatternFill("solid", fgColor="C8004B")
_POLICE_TOTAL = Font(bold=True)


def _anonymiser(profil):
    """Identifiant technique stable, jamais le nom."""
    return f"ETU-{profil.pk:05d}"


def _ecrire_entetes(feuille, ligne, colonnes):
    for i, texte in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=ligne, column=i, value=texte)
        cellule.font = _POLICE_ENTETE
        cellule.fill = _REMPLISSAGE_ENTETE


def _largeurs(feuille, largeurs):
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = largeur


def _type_transaction(tx):
    """Classe une transaction pour l'affichage : Evenement (avec son nom),
    Vente libre, ou Produit. Une transaction qui melangerait plusieurs
    types (rare, ex. un panier avec une place ET un cafe) est classee sur
    le premier type rencontre, evenement d'abord."""
    lignes = list(tx.lignes.all())
    for ligne in lignes:
        if ligne.produit.evenement_id:
            return f"Evenement : {ligne.produit.evenement.nom}"
    for ligne in lignes:
        if ligne.produit.est_vente_libre:
            return "Vente libre"
    return "Produit"


def construire_classeur(pole, date_debut, date_fin):
    """date_debut incluse, date_fin EXCLUE (bornes demi-ouvertes), toutes
    deux des datetime ou date compatibles avec un filtre Django."""
    transactions = (
        Transaction.objects.filter(
            pole=pole, date_operation__gte=date_debut, date_operation__lt=date_fin
        )
        .select_related("profil")
        .prefetch_related("lignes__produit__evenement")
        .order_by("date_operation")
    )
    lignes = list(
        LigneTransaction.objects.filter(transaction__in=transactions)
        .select_related("produit__evenement", "transaction__profil")
    )

    recette_evenements = Decimal("0")
    recette_produits = Decimal("0")
    recette_autre = Decimal("0")
    for ligne in lignes:
        montant = ligne.prix_unitaire * ligne.quantite
        if ligne.produit.evenement_id:
            recette_evenements += montant
        elif ligne.produit.est_vente_libre:
            recette_autre += montant
        else:
            recette_produits += montant
    recette_totale = recette_evenements + recette_produits + recette_autre

    classeur = Workbook()

    # ------------------------------------------------------------------
    # Onglet 1 : Transactions (liste + recapitulatif a cote)
    # ------------------------------------------------------------------
    feuille = classeur.active
    feuille.title = "Transactions"
    _ecrire_entetes(feuille, 1, ["Date", "Acheteur", "Type", "Montant (EUR)"])
    row = 2
    for tx in transactions:
        feuille.cell(row=row, column=1, value=tx.date_operation.strftime("%d/%m/%Y %H:%M"))
        feuille.cell(row=row, column=2, value=_anonymiser(tx.profil))
        feuille.cell(row=row, column=3, value=_type_transaction(tx))
        feuille.cell(row=row, column=4, value=float(tx.montant_total))
        row += 1
    if row == 2:
        feuille.cell(row=2, column=1, value="Aucune transaction sur cette periode.")
    _largeurs(feuille, [18, 14, 30, 16])

    # Tableau recapitulatif, decale a droite pour ne pas gener la liste.
    col = 6
    for i, texte in enumerate(["Recapitulatif", "Montant (EUR)"]):
        cellule = feuille.cell(row=1, column=col + i, value=texte)
        cellule.font = _POLICE_ENTETE
        cellule.fill = _REMPLISSAGE_ENTETE
    lignes_recap = [
        ("Recette evenements", recette_evenements),
        ("Recette produits", recette_produits),
        ("Recette vente libre", recette_autre),
        ("Total", recette_totale),
    ]
    for i, (libelle, montant) in enumerate(lignes_recap, start=2):
        c1 = feuille.cell(row=i, column=col, value=libelle)
        c2 = feuille.cell(row=i, column=col + 1, value=float(montant))
        if libelle == "Total":
            c1.font = _POLICE_TOTAL
            c2.font = _POLICE_TOTAL
    feuille.column_dimensions[get_column_letter(col)].width = 22
    feuille.column_dimensions[get_column_letter(col + 1)].width = 16

    # ------------------------------------------------------------------
    # Onglet 2 : Evenements (un tableau distinct par evenement)
    # ------------------------------------------------------------------
    feuille_evt = classeur.create_sheet("Evenements")
    lignes_par_evenement = defaultdict(list)
    for ligne in lignes:
        if ligne.produit.evenement_id:
            lignes_par_evenement[ligne.produit.evenement].append(ligne)

    row = 1
    if not lignes_par_evenement:
        feuille_evt.cell(row=1, column=1, value="Aucune vente d'evenement sur cette periode.")
    for evenement, lignes_evt in lignes_par_evenement.items():
        feuille_evt.cell(row=row, column=1, value=evenement.nom).font = Font(bold=True, size=13)
        row += 1
        _ecrire_entetes(feuille_evt, row, ["Date", "Acheteur", "Produit", "Quantite", "Montant (EUR)"])
        row += 1
        total_evt = Decimal("0")
        for ligne in sorted(lignes_evt, key=lambda l: l.transaction.date_operation):
            montant = ligne.prix_unitaire * ligne.quantite
            total_evt += montant
            feuille_evt.cell(row=row, column=1, value=ligne.transaction.date_operation.strftime("%d/%m/%Y %H:%M"))
            feuille_evt.cell(row=row, column=2, value=_anonymiser(ligne.transaction.profil))
            feuille_evt.cell(row=row, column=3, value=ligne.libelle)
            feuille_evt.cell(row=row, column=4, value=ligne.quantite)
            feuille_evt.cell(row=row, column=5, value=float(montant))
            row += 1
        c1 = feuille_evt.cell(row=row, column=4, value="Total")
        c2 = feuille_evt.cell(row=row, column=5, value=float(total_evt))
        c1.font = _POLICE_TOTAL
        c2.font = _POLICE_TOTAL
        row += 2  # ligne vide entre deux evenements
    _largeurs(feuille_evt, [18, 14, 28, 10, 16])

    # ------------------------------------------------------------------
    # Onglet 3 : Produits (catalogue normal, hors evenements/terminal)
    # ------------------------------------------------------------------
    feuille_prod = classeur.create_sheet("Produits")
    lignes_par_produit = defaultdict(list)
    for ligne in lignes:
        if not ligne.produit.evenement_id and not ligne.produit.est_vente_libre:
            lignes_par_produit[ligne.produit.nom].append(ligne)

    row = 1
    _ecrire_entetes(
        feuille_prod, row,
        ["Date", "Acheteur", "Produit", "Quantite", "Prix unitaire (EUR)", "Montant (EUR)"],
    )
    row += 1
    if not lignes_par_produit:
        feuille_prod.cell(row=row, column=1, value="Aucune vente de produit sur cette periode.")
        row += 1
    total_general = Decimal("0")
    for nom_produit in sorted(lignes_par_produit):
        lignes_p = lignes_par_produit[nom_produit]
        sous_total = Decimal("0")
        for ligne in sorted(lignes_p, key=lambda l: l.transaction.date_operation):
            montant = ligne.prix_unitaire * ligne.quantite
            sous_total += montant
            feuille_prod.cell(row=row, column=1, value=ligne.transaction.date_operation.strftime("%d/%m/%Y %H:%M"))
            feuille_prod.cell(row=row, column=2, value=_anonymiser(ligne.transaction.profil))
            feuille_prod.cell(row=row, column=3, value=ligne.libelle)
            feuille_prod.cell(row=row, column=4, value=ligne.quantite)
            feuille_prod.cell(row=row, column=5, value=float(ligne.prix_unitaire))
            feuille_prod.cell(row=row, column=6, value=float(montant))
            row += 1
        c1 = feuille_prod.cell(row=row, column=3, value=f"Sous-total {nom_produit}")
        c2 = feuille_prod.cell(row=row, column=6, value=float(sous_total))
        c1.font = _POLICE_TOTAL
        c2.font = _POLICE_TOTAL
        total_general += sous_total
        row += 2
    c1 = feuille_prod.cell(row=row, column=3, value="Total general")
    c2 = feuille_prod.cell(row=row, column=6, value=float(total_general))
    c1.font = _POLICE_TOTAL
    c2.font = _POLICE_TOTAL
    _largeurs(feuille_prod, [18, 14, 28, 10, 16, 16])

    return classeur
